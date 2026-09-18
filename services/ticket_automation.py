"""Tickets Automation — Jira admin-note tickets enriched with session-log URLs,
exportable to Excel (ported from Jira-Ticket-Automation)."""
import io
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import jira_common
from .mongo_common import get_collection

SESSION_URL_BASE = "https://platform.dxfactor.com/member-concierge/chat-log"

COLUMNS = ["Ticket ID", "Session ID", "Session URL", "Note", "Root Cause", "Recommended Fix", "Fixed"]
COL_WIDTHS = [15, 35, 60, 50, 30, 30, 12]
HEADER_BG = "1E3A5F"

SESSION_ID_RE = re.compile(r"(session_\S+)")
DIRECT_DOC_ID_RE = re.compile(r"chat-log\?id=([0-9a-fA-F]{24})")

BLOCK_NODE_TYPES = (
    "doc", "paragraph", "heading", "bulletList", "orderedList", "listItem",
    "table", "tableRow", "tableCell", "tableHeader", "panel", "blockquote", "rule",
)


def _extract_text_from_adf(node) -> str:
    if isinstance(node, str):
        return node
    if not isinstance(node, dict):
        return ""
    if node.get("type") == "text":
        return node.get("text", "")
    if node.get("type") == "hardBreak":
        return "\n"
    parts = [_extract_text_from_adf(child) for child in node.get("content", [])]
    separator = "\n" if node.get("type") in BLOCK_NODE_TYPES else ""
    return separator.join(p for p in parts if p)


def _parse_description(raw, summary: str = "") -> dict:
    text = _extract_text_from_adf(raw) if isinstance(raw, dict) else (raw or "")
    session = re.search(r"Conversation ID:\s*(\S+)", text)
    if not session:
        session = SESSION_ID_RE.search(summary or "")
    org = re.search(r"Organization ID:\s*(\d+)", text)
    org_name = re.search(r"Organization:\s*(.+)", text)
    agent = re.search(r"Agent:\s*(.+)", text)
    note_match = re.search(r"Admin Note\s*\n+(.*?)\n*(?:Conversation Logs|$)", text, re.DOTALL | re.IGNORECASE)
    if not note_match:
        note_match = re.search(r"Note Content\s*\n+(.*?)\n*(?:Conversation Link|Conversation Logs|$)", text, re.DOTALL | re.IGNORECASE)
    direct_doc = DIRECT_DOC_ID_RE.search(text)
    return {
        "session_id": session.group(1).strip() if session else "",
        "org_id": org.group(1).strip() if org else "",
        "org_name": org_name.group(1).strip() if org_name else "",
        "agent_name": agent.group(1).strip() if agent else "",
        "note": note_match.group(1).strip() if note_match else "",
        "direct_doc_id": direct_doc.group(1) if direct_doc else "",
    }


def _org_or_agent_matches(org_id: str, org_name: str, agent_name: str, parsed: dict) -> bool:
    if org_name and parsed.get("org_name", "").strip().lower() != org_name.strip().lower():
        return False
    if parsed.get("org_id"):
        return not org_id or parsed["org_id"] == org_id
    if agent_name:
        return parsed.get("agent_name", "").strip().lower() == agent_name.strip().lower()
    return not org_id and not org_name


def _matches(ticket: dict, org_id: str, org_name: str, agent_name: str, date_from, date_to, parsed: dict) -> bool:
    if not _org_or_agent_matches(org_id, org_name, agent_name, parsed):
        return False
    if date_from or date_to:
        created_str = ticket.get("fields", {}).get("created", "")
        try:
            created = datetime.fromisoformat(created_str.replace("Z", "+00:00")).date()
        except Exception:
            return False
        if date_from and created < date_from:
            return False
        if date_to and created > date_to:
            return False
    return True


def _get_document_ids(session_ids: list[str]) -> dict[str, str]:
    if not session_ids:
        return {}
    docs = get_collection().find({"session_id": {"$in": session_ids}}, {"_id": 1, "session_id": 1})
    return {doc["session_id"]: str(doc["_id"]) for doc in docs}


def process(org_id: str, agent_name: str = "", org_name: str = "", date_from=None, date_to=None) -> dict:
    cfg = jira_common.get_config()
    clauses = []
    if org_id:
        clauses.append(f'description ~ "Organization ID: {org_id}"')
    if org_name:
        clauses.append(f'description ~ "Organization: {org_name}"')
    if agent_name:
        clauses.append(f'description ~ "Agent: {agent_name}"')
    filter_clause = f" AND ({' OR '.join(clauses)})" if clauses else ""
    jql = f'project = {cfg["project"]}{filter_clause} ORDER BY created DESC'
    tickets = jira_common.search_jql(cfg, jql, ["summary", "description", "status", "created"])

    matched = []
    for ticket in tickets:
        fields = ticket.get("fields", {})
        parsed = _parse_description(fields.get("description") or "", fields.get("summary") or "")
        if _matches(ticket, org_id, org_name, agent_name, date_from, date_to, parsed):
            matched.append((ticket, parsed))

    session_ids = [p["session_id"] for _, p in matched if p["session_id"] and not p["direct_doc_id"]]
    doc_id_map = _get_document_ids(session_ids) if session_ids else {}

    rows = []
    for ticket, parsed in matched:
        session_id = parsed["session_id"]
        document_id = parsed["direct_doc_id"] or doc_id_map.get(session_id)
        fields = ticket.get("fields", {})
        rows.append({
            "ticket_id": ticket.get("key", ""),
            "session_id": session_id,
            "session_url": f"{SESSION_URL_BASE}/{document_id}" if document_id else "",
            "agent_name": parsed["agent_name"],
            "organization": parsed["org_name"] or org_name or org_id,
            "note": parsed["note"],
            "status": (fields.get("status") or {}).get("name", ""),
            "created": (fields.get("created") or "")[:10],
        })

    return {
        "org_id": org_id,
        "org_name": org_name,
        "agent_name": agent_name,
        "fetched": len(tickets),
        "matched": len(rows),
        "resolved": len([r for r in rows if r["session_url"]]),
        "unresolved": len([r for r in rows if not r["session_url"]]),
        "rows": rows,
    }


def generate_excel(org_id: str, rows: list[dict], org_name: str = "") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(org_id)

    ws.append([f"Organization: {org_name or org_id}"])
    ws["A1"].font = Font(bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    ws.append(COLUMNS)
    fill = PatternFill("solid", fgColor=HEADER_BG)
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[2]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[ws.cell(2, i).column_letter].width = w

    for row in rows:
        ws.append([
            row.get("ticket_id", ""),
            row.get("session_id", ""),
            row.get("session_url", ""),
            row.get("note", ""),
            "", "", "",
        ])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
